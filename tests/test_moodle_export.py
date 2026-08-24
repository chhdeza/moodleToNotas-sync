"""
Pruebas de la descarga de notas desde Moodle.

Las dos que más importan no son las del camino feliz, sino las dos rejas:
que no se pueda descargar el curso entero creyendo que es un grupo, y que no
se pueda seguir adelante sin las columnas que identifican al estudiante.
"""

from __future__ import annotations

import pytest
import responses

from mnsync.errors import MoodleError, ScopeError
from mnsync.moodle_export import (
    COL_CEDULA,
    COL_INSTITUCION,
    MoodleSession,
    is_grade_column,
    write_moodle_xlsx,
)
from tests.fixtures import moodle_html as fx
from tests.fixtures.gen import make_students

BASE = "http://127.0.0.1:9/moodle"
CURSO = 8067
GRUPO = 38525


def _mock_login(rsps):
    rsps.add(responses.GET, f"{BASE}/login/index.php", body=fx.LOGIN_PAGE, status=200)
    rsps.add(responses.POST, f"{BASE}/login/index.php", body=fx.LOGIN_OK, status=200)


def _session():
    s = MoodleSession(BASE)
    s.login("profesor.prueba", "contrasena-de-prueba")
    return s


# ---------------------------------------------------------------------------
# Ingreso
# ---------------------------------------------------------------------------
@responses.activate
def test_login_ok():
    _mock_login(responses.mock)
    _session()  # no debe lanzar


@responses.activate
def test_login_rechazado_dice_que_revisar():
    responses.add(responses.GET, f"{BASE}/login/index.php", body=fx.LOGIN_PAGE, status=200)
    responses.add(responses.POST, f"{BASE}/login/index.php", body=fx.LOGIN_FAIL, status=200)

    s = MoodleSession(BASE)
    with pytest.raises(MoodleError) as ex:
        s.login("profesor.prueba", "mala")
    assert "MOODLE_USERNAME" in str(ex.value)


# ---------------------------------------------------------------------------
# Grupos
# ---------------------------------------------------------------------------
@responses.activate
def test_list_groups_omite_todos_los_participantes():
    _mock_login(responses.mock)
    responses.add(
        responses.GET,
        f"{BASE}/grade/export/txt/index.php",
        body=fx.export_page(course_id=CURSO),
        status=200,
    )

    grupos = _session().list_groups(CURSO)
    assert [(g.id, g.name) for g in grupos] == [(38525, "Grupo 1"), (38526, "Grupo 2")]


# ---------------------------------------------------------------------------
# La reja de alcance
# ---------------------------------------------------------------------------
@responses.activate
def test_aborta_si_moodle_ignora_el_filtro_de_grupo():
    """
    El fallo peligroso: se pide un grupo y Moodle sirve el curso entero.

    No da error HTTP, así que hay que detectarlo mirando qué quedó
    seleccionado. Sin esta reja se subirían notas de estudiantes ajenos.
    """
    _mock_login(responses.mock)
    responses.add(
        responses.GET,
        f"{BASE}/grade/export/txt/index.php",
        body=fx.export_page(course_id=CURSO, selected_group=None),
        status=200,
    )

    with pytest.raises(ScopeError) as ex:
        _session().export_group(CURSO, GRUPO)
    assert "38525" in str(ex.value)
    assert "courses.yml" in str(ex.value)


@responses.activate
def test_aborta_si_el_curso_no_tiene_grupos():
    _mock_login(responses.mock)
    responses.add(
        responses.GET,
        f"{BASE}/grade/export/txt/index.php",
        body=fx.export_page(course_id=CURSO, with_group_select=False),
        status=200,
    )

    with pytest.raises(ScopeError) as ex:
        _session().export_group(CURSO, GRUPO)
    assert "otros profesores" in str(ex.value)


@responses.activate
def test_aborta_si_moodle_selecciona_otro_grupo():
    _mock_login(responses.mock)
    responses.add(
        responses.GET,
        f"{BASE}/grade/export/txt/index.php",
        body=fx.export_page(course_id=CURSO, selected_group=38526),
        status=200,
    )

    with pytest.raises(ScopeError):
        _session().export_group(CURSO, 38525)


# ---------------------------------------------------------------------------
# La reja de columnas
# ---------------------------------------------------------------------------
@responses.activate
def test_aborta_sin_columna_de_cedula_y_explica_como_arreglarlo():
    _mock_login(responses.mock)
    responses.add(
        responses.GET,
        f"{BASE}/grade/export/txt/index.php",
        body=fx.export_page(course_id=CURSO, selected_group=GRUPO),
        status=200,
    )
    responses.add(
        responses.POST,
        f"{BASE}/grade/export/txt/export.php",
        body="Nombre,Apellido(s),Tarea 1 (Real)\nAna,Mora,89\n",
        status=200,
    )

    with pytest.raises(MoodleError) as ex:
        _session().export_group(CURSO, GRUPO)
    msg = str(ex.value)
    assert COL_CEDULA in msg and COL_INSTITUCION in msg
    # El error tiene que decir a quién pedirle el cambio, no solo que falta.
    assert "administra Moodle" in msg


# ---------------------------------------------------------------------------
# Camino feliz
# ---------------------------------------------------------------------------
@responses.activate
def test_export_group_marca_todos_los_instrumentos_y_fuerza_nota_real():
    _mock_login(responses.mock)
    responses.add(
        responses.GET,
        f"{BASE}/grade/export/txt/index.php",
        body=fx.export_page(course_id=CURSO, selected_group=GRUPO),
        status=200,
    )
    estudiantes = make_students(3, cu="42")
    csv_body = fx.export_csv(
        estudiantes,
        {estudiantes[0].cedula: {"Tarea 1 (Real)": 89, "Proyecto (Real)": 95}},
    )
    responses.add(responses.POST, f"{BASE}/grade/export/txt/export.php", body=csv_body, status=200)

    export = _session().export_group(CURSO, GRUPO)

    assert len(export) == 3
    assert export.cedulas == {s.cedula for s in estudiantes}
    assert export.grade_headers == ["Tarea 1 (Real)", "Proyecto (Real)"]

    enviado = responses.calls[-1].request.body
    # Las tres casillas de instrumento, incluida la que venia desmarcada.
    for item in (101, 102, 103):
        assert f"itemids%5B{item}%5D=1" in enviado
    # Nota real, sin porcentaje ni comentarios.
    assert "display%5Breal%5D=1" in enviado
    assert "display%5Bpercentage%5D" not in enviado
    assert "export_feedback=0" in enviado
    # El anti-CSRF cosechado del formulario viaja de vuelta.
    assert "sesskey=s3sskeyPrueba" in enviado


@responses.activate
def test_html_en_vez_de_csv_se_reporta_como_sesion_vencida():
    _mock_login(responses.mock)
    responses.add(
        responses.GET,
        f"{BASE}/grade/export/txt/index.php",
        body=fx.export_page(course_id=CURSO, selected_group=GRUPO),
        status=200,
    )
    responses.add(
        responses.POST,
        f"{BASE}/grade/export/txt/export.php",
        body="<html><body>Su sesion ha expirado</body></html>",
        status=200,
    )

    with pytest.raises(MoodleError) as ex:
        _session().export_group(CURSO, GRUPO)
    assert "--from-xlsx" in str(ex.value)


# ---------------------------------------------------------------------------
# Columnas de nota y escritura del xlsx
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "header,esperado",
    [
        ("Tarea 1 (Real)", True),
        ("Proyecto Final", True),  # export simple, sin sufijo
        ("Nombre", False),
        ("Número de ID", False),
        ("Institución", False),
        ("Correo electrónico", False),
        ("", False),
    ],
)
def test_is_grade_column(header, esperado):
    assert is_grade_column(header) is esperado


@responses.activate
def test_write_moodle_xlsx_pone_los_encabezados_que_espera_el_uploader(tmp_path):
    _mock_login(responses.mock)
    responses.add(
        responses.GET,
        f"{BASE}/grade/export/txt/index.php",
        body=fx.export_page(course_id=CURSO, selected_group=GRUPO),
        status=200,
    )
    estudiantes = make_students(2, cu="42")
    responses.add(
        responses.POST,
        f"{BASE}/grade/export/txt/export.php",
        body=fx.export_csv(estudiantes, {estudiantes[0].cedula: {"Tarea 1 (Real)": 89}}),
        status=200,
    )

    export = _session().export_group(CURSO, GRUPO)
    destino = write_moodle_xlsx(export, tmp_path / "calificaciones.xlsx")

    import openpyxl

    ws = openpyxl.load_workbook(destino).active
    encabezados = [c.value for c in ws[1]]
    # Los cuatro nombres exactos que busca _ingest_xlsx_files.
    assert encabezados[:4] == ["Nombre", "Apellido(s)", "Número de ID", "Institución"]
    assert "Tarea 1 (Real)" in encabezados

    filas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(filas) == 2
    # El guion de "sin calificar" se preserva tal cual: el planificador lo
    # interpreta como "no presento".
    assert "-" in [str(v) for v in filas[1]]


# ---------------------------------------------------------------------------
# Columnas que Moodle calcula solo
# ---------------------------------------------------------------------------
def test_el_total_del_curso_no_es_una_nota():
    """
    Moodle agrega «Total del curso» solo, y no es un instrumento.

    Notas Parciales saca su propio promedio de las notas cargadas, así que esa
    columna no tiene dónde subirse. Si se dejara pasar, cada corrida terminaría
    con una advertencia por estudiante diciendo que no se pudo emparejar — y una
    advertencia que sale siempre y nunca significa nada enseña a no leerlas.
    """
    from mnsync.moodle_export import es_columna_calculada, is_grade_column

    for calculada in (
        "Total del curso (Real)",
        "Total del curso",
        "Total de categoría (Real)",
        "Total de la categoría (Real)",
    ):
        assert es_columna_calculada(calculada), calculada
        assert not is_grade_column(calculada), calculada


def test_una_tarea_de_verdad_sigue_siendo_una_nota():
    """El filtro no puede llevarse por delante una columna calificable."""
    from mnsync.moodle_export import is_grade_column

    for real in (
        "Tarea:Entrega de Actividad Tarea 1 (Real)",
        "Proyecto (Real)",
        "Total de puntos extra",  # empieza con «Total» pero no es un agregado
    ):
        assert is_grade_column(real), real


def test_el_xlsx_no_lleva_correos_ni_totales(tmp_path):
    """
    Al script solo se le pasa lo que necesita: identidad y notas.

    Además de quitar ruido, deja de escribir el correo de cada estudiante en un
    archivo del disco, que no hacía falta para nada.
    """
    import openpyxl

    from mnsync.moodle_export import GradeExport, write_moodle_xlsx

    export = GradeExport(
        headers=[
            "Nombre", "Apellido(s)", "Número de ID", "Institución",
            "Dirección de correo", "Tarea 1 (Real)", "Total del curso (Real)",
        ],
        rows=[
            {
                "Nombre": "ANA", "Apellido(s)": "SOLANO",
                "Número de ID": "0100000001", "Institución": "DESAMPARADOS (42)",
                "Dirección de correo": "ana@example.com",
                "Tarea 1 (Real)": "80", "Total del curso (Real)": "80",
            }
        ],
        group_id=1,
        grade_headers=["Tarea 1 (Real)"],
    )

    ruta = write_moodle_xlsx(export, tmp_path / "x.xlsx")
    wb = openpyxl.load_workbook(ruta, read_only=True)
    encabezados = [str(c or "") for c in next(wb.active.iter_rows(values_only=True))]
    wb.close()

    assert encabezados == [
        "Nombre", "Apellido(s)", "Número de ID", "Institución", "Tarea 1 (Real)",
    ]


def test_item_map_puede_rescatar_una_columna_excluida(tmp_path):
    """
    Si el profesor dice a mano que esa columna va a un instrumento, va.

    El filtro es un valor por defecto sensato, no una decisión sobre la que él
    no pueda mandar.
    """
    import openpyxl

    from mnsync.moodle_export import GradeExport, write_moodle_xlsx

    export = GradeExport(
        headers=["Nombre", "Apellido(s)", "Número de ID", "Institución", "Total del curso (Real)"],
        rows=[],
        group_id=1,
        grade_headers=[],
    )

    ruta = write_moodle_xlsx(
        export, tmp_path / "x.xlsx", conservar={"Total del curso (Real)"}
    )
    wb = openpyxl.load_workbook(ruta, read_only=True)
    encabezados = [str(c or "") for c in next(wb.active.iter_rows(values_only=True))]
    wb.close()

    assert "Total del curso (Real)" in encabezados
